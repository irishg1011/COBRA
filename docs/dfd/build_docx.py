"""Builds CobraByte_Context_Diagram_and_DFD.docx from png/*.png, specs.js and descriptions.js.

Laid out like the capstone papers it is modelled on: portrait Letter, 1 in
margins, Arial 12. Each figure starts a page: the diagram centred, the caption
("Figure N" / title) centred under it, then the description - double-spaced,
justified, first line indented 0.5 in.
"""
import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / 'CobraByte_Context_Diagram_and_DFD.docx'
FIRST_FIGURE = int(sys.argv[2]) if len(sys.argv) > 2 else 1  # change to continue your paper's figure numbers

data = json.loads(subprocess.check_output(
    ['node', '-e', 'console.log(JSON.stringify({specs: require("./specs.js"), text: require("./descriptions.js")}))'], cwd=HERE))
specs, text = data['specs'], data['text']
meta = {m['id']: m for m in json.load(open(HERE / 'png' / 'meta.json'))}

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
    setattr(sec, side, Inches(1))
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
normal.font.size = Pt(12)


def spacing(p, double=True):
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(0), Pt(0)
    f.line_spacing_rule = WD_LINE_SPACING.DOUBLE if double else WD_LINE_SPACING.SINGLE


MAX_W, MAX_H = 6.5, 7.6       # room for the caption under the diagram on the same page
MAX_SCALE = 0.01              # in per px: small diagrams keep about 9 pt text instead of filling the page
LEVEL_ORDER = ['Context', '0', '1', '2', '3', '4']
diagrams = sorted((d for d in specs if d['level'] != 'legend'), key=lambda d: LEVEL_ORDER.index(d['level']))
missing = [d['id'] for d in diagrams if d['id'] not in text]
if missing:
    sys.exit('no description for: ' + ', '.join(missing))

for i, d in enumerate(diagrams):
    m = meta[d['id']]
    scale = min(MAX_W / m['w'], MAX_H / m['h'], MAX_SCALE)
    pic = doc.add_paragraph()
    spacing(pic, double=False)
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.page_break_before = i > 0
    pic.paragraph_format.keep_with_next = True
    pic.add_run().add_picture(str(HERE / 'png' / (d['id'] + '.png')), width=Inches(m['w'] * scale))

    cap = doc.add_paragraph()
    spacing(cap)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(f'Figure {FIRST_FIGURE + i}')
    run.add_break(WD_BREAK.LINE)
    cap.add_run(d['title'])

    for para in text[d['id']]:
        p = doc.add_paragraph(para)
        spacing(p)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0.5)

doc.save(OUT)
print('saved', OUT, len(diagrams), 'figures')
